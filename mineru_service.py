# -*- coding: utf-8 -*-
import os
import uuid
import shutil
import time
import json
import asyncio
import httpx
import logging
from pathlib import Path
from typing import List, Optional
from pydantic import BaseModel
from fastapi import FastAPI, Query, UploadFile, File, Form, HTTPException
from fastapi.responses import JSONResponse, RedirectResponse
import markdown
import uvicorn
import subprocess

# 依赖
import pandas as pd
from tabulate import tabulate
from docx import Document

# 设置日志
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger("mineru")

# 创建 FastAPI 应用
app = FastAPI()

# 临时输出目录
TEMP_BASE_DIR = Path("./temp_output")
TEMP_BASE_DIR.mkdir(exist_ok=True)


# 启动清理任务（清除超过 1 小时的目录）
async def clean_temp_dirs():
    while True:
        try:
            for dir in TEMP_BASE_DIR.iterdir():
                if dir.is_dir() and (time.time() - dir.stat().st_mtime) > 3600:
                    shutil.rmtree(dir, ignore_errors=True)
                    logger.info(f"清理目录: {dir}")
        except Exception as e:
            logger.exception(f"清理任务异常: {e}")
        await asyncio.sleep(3600)


@app.on_event("startup")
async def startup_event():
    asyncio.create_task(clean_temp_dirs())


# 请求体模型
class ParseRequest(BaseModel):
    file_paths: Optional[List[str]] = None
    lang: Optional[str] = "ch"
    backend: Optional[str] = "pipeline"
    method: Optional[str] = "auto"
    output_format: Optional[str] = "md"


# Word (.docx) 转 Markdown
def docx_to_md(input_path: Path) -> str:
    doc = Document(input_path)
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if para.style.name.startswith("Heading"):
            level = para.style.name.replace("Heading", "").strip()
            if level.isdigit():
                lines.append("#" * int(level) + " " + text)
            else:
                lines.append(text)
        else:
            lines.append(text)
    for table in doc.tables:
        data = []
        for row in table.rows:
            data.append([cell.text.strip() for cell in row.cells])
        lines.append(tabulate(data, headers="firstrow", tablefmt="github"))
    return "\n\n".join(lines)


# 兼容老 .doc（需要 antiword 工具）
def doc_to_md(input_path: Path) -> str:
    try:
        result = subprocess.run(
            ["antiword", str(input_path)],
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            timeout=30,
            check=True
        )
        text = result.stdout.decode("utf-8", errors="ignore")
        return text
    except FileNotFoundError:
        raise RuntimeError("不是标准 .docx 压缩包，且未找到可用的外部解析工具（antiword）。请安装 antiword 或将 .doc 另存为 .docx。")
    except subprocess.CalledProcessError as e:
        raise RuntimeError(f".doc 文件解析失败: {e.stderr.decode('utf-8', errors='ignore')}")


# Excel 转 Markdown
def excel_to_md(input_path: Path) -> str:
    excel_file = pd.ExcelFile(input_path)
    md_parts = []
    for sheet in excel_file.sheet_names:
        df = excel_file.parse(sheet)
        md_table = tabulate(df.fillna(""), headers="keys", tablefmt="github", showindex=False)
        md_parts.append(f"## {sheet}\n\n{md_table}")
    return "\n\n".join(md_parts)


# 调用 mineru 命令行接口（只处理 PDF）
def run_mineru_cli(input_path: str, output_path: str, lang: str = "ch", backend: str = "pipeline", method: str = "auto"):
    os.environ['MINERU_MODEL_SOURCE'] = "modelscope"
    cmd = f"mineru -p {input_path} -o {output_path} --source local"
    logger.info(f"运行命令: {cmd}")
    code = os.system(cmd)
    if code != 0:
        raise RuntimeError(f"mineru 命令执行失败，退出码 {code}")


@app.post("/parse")
async def parse_from_path(req: ParseRequest):
    result = {}
    for path in req.file_paths:
        suffix = Path(path).suffix.lower()
        name = Path(path).stem

        if suffix == ".pdf":
            # PDF 调用 mineru
            task_id = uuid.uuid4().hex[:8]
            task_output_dir = TEMP_BASE_DIR / task_id
            task_output_dir.mkdir(exist_ok=True)
            run_mineru_cli(path, str(task_output_dir), lang=req.lang, backend=req.backend, method=req.method)
            output_md_path = task_output_dir / name / "auto" / f"{name}.md"
            md_content = output_md_path.read_text(encoding="utf-8")
        elif suffix == ".docx":
            md_content = docx_to_md(Path(path))
        elif suffix == ".doc":
            md_content = doc_to_md(Path(path))
        elif suffix in [".xls", ".xlsx"]:
            md_content = excel_to_md(Path(path))
        else:
            raise HTTPException(status_code=400, detail=f"不支持的文件类型: {suffix}")

        if req.output_format == "html":
            result[name] = convert_md_to_html(md_content)
        else:
            result[name] = md_content
    return JSONResponse(content=result)


@app.post("/upload")
async def parse_from_upload(
    file: UploadFile = File(...),
    lang: str = Form("ch"),
    backend: str = Form("pipeline"),
    method: str = Form("auto"),
    output_format: str = Form("md")
):
    content = await file.read()
    name = Path(file.filename).stem
    suffix = Path(file.filename).suffix.lower()

    task_id = uuid.uuid4().hex[:8]
    task_dir = TEMP_BASE_DIR / task_id
    task_dir.mkdir(exist_ok=True)

    input_path = task_dir / file.filename
    with open(input_path, "wb") as f:
        f.write(content)

    if suffix == ".pdf":
        run_mineru_cli(str(input_path), str(task_dir), lang=lang, backend=backend, method=method)
        output_md_path = task_dir / name / "auto" / f"{name}.md"
        md_content = output_md_path.read_text(encoding="utf-8")
    elif suffix == ".docx":
        md_content = docx_to_md(input_path)
    elif suffix == ".doc":
        md_content = doc_to_md(input_path)
    elif suffix in [".xls", ".xlsx"]:
        md_content = excel_to_md(input_path)
    else:
        raise HTTPException(status_code=400, detail=f"不支持的文件类型: {suffix}")

    if output_format == "html":
        return convert_md_to_html(md_content)
    else:
        return md_content


def convert_md_to_html(md_text: str) -> str:
    html = markdown.markdown(md_text, extensions=['extra', 'tables', 'nl2br'])
    html = html.replace("\n", "")
    return html


# 远程文件服务器地址列表
servers = [
    'http://122.9.181.18:89/'
]


# 检查远程文件是否存在
async def check_file_exists(filename: str):
    async with httpx.AsyncClient() as client:
        for i, base_url in enumerate(servers):
            file_url = base_url + filename
            try:
                resp = await client.head(file_url, timeout=5)
                if resp.status_code == 200:
                    return {"file_url": file_url, "server_index": i}
            except Exception as e:
                logger.warning(f"Error checking {file_url}: {e}")
    return None


@app.get("/file")
async def handle_file(
    filename: str = Query(..., description="要检查或下载的文件名"),
    action: str = Query("check", description="操作类型：check 或 download")
):
    if not filename:
        return JSONResponse(content={"error": "filename is required"}, status_code=400)

    file_data = await check_file_exists(filename)
    if not file_data:
        return JSONResponse(content={"exists": False}, status_code=404)

    file_url = f"{servers[0]}{filename}"

    if action == "check":
        return JSONResponse(content={"exists": True, "serverIndex": file_data["server_index"]})
    elif action == "download":
        return RedirectResponse(url=file_url, status_code=307)
    else:
        return JSONResponse(content={"error": "Invalid action"}, status_code=400)


# 启动服务
if __name__ == '__main__':
    uvicorn.run("__main__:app", host="0.0.0.0", port=5002, reload=False)
